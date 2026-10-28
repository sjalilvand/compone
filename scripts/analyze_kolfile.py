#!/usr/bin/env python3
"""Analyze kolfile.xlsx structure: merged cells, columns, multi-level headers, RTL."""
import openpyxl

path = '/home/z/my-project/upload/kolfile.xlsx'
wb = openpyxl.load_workbook(path)

print("=== Sheet names ===")
print(wb.sheetnames)

for ws in wb.worksheets:
    print(f"\n=== Sheet: {ws.title} ===")
    print(f"Dimensions: {ws.dimensions}, max_row={ws.max_row}, max_col={ws.max_column}")
    print(f"RTL (rightToLeft): {ws.sheet_view.rightToLeft}")
    print(f"Merged cells: {list(ws.merged_cells.ranges)}")
    print("\n--- Cell contents (first 12 rows) ---")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12)):
        for cell in row:
            if cell.value is not None:
                v = str(cell.value).replace('\n', ' | ')
                print(f"  {cell.coordinate}: {repr(v)}")
    print("\n--- Column widths ---")
    for letter, dim in ws.column_dimensions.items():
        if dim.width:
            print(f"  {letter}: {dim.width}")
    print("\n--- Fonts of first 6 rows (non-empty, unique) ---")
    seen = set()
    for row in ws.iter_rows(min_row=1, max_row=6):
        for cell in row:
            if cell.value is not None:
                f = cell.font
                key = (f.name, f.size, f.bold)
                if key not in seen:
                    seen.add(key)
                    align = cell.alignment
                    print(f"  {cell.coordinate}: font={f.name}, size={f.size}, bold={f.bold}, halign={align.horizontal}, wrap={align.wrap_text}")
    print("\n--- Any formulas/data validations ---")
    print(f"  Data validations: {len(ws.data_validations.dataValidation) if ws.data_validations else 0}")
